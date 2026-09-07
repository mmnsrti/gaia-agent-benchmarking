"""Unit tests for V2 File / Attachment handling tool, prompts, agent, and runner.

All tests use local files or mocks to ensure no external network or API calls are made.
"""

import os
import tempfile
import unittest
from unittest.mock import patch, MagicMock

import docx
import openpyxl
from pptx import Presentation
from pptx.util import Inches

from tools.file_tool import FileTool, FileResult
from tools.web_search import WebSearchResult, SearchResultItem
from agent.agent import GAIAFileAgent, GAIAWebAgent, GAIAAgent, AgentResult
from agent.llm import LLMResponse
from prompts.file_search import build_file_search_prompt, FILE_SEARCH_PROMPT_VERSION
from prompts.web_search import WEB_SEARCH_PROMPT_VERSION
from prompts.baseline import PROMPT_VERSION as BASELINE_PROMPT_VERSION
from evaluation.runner import execute_task
from evaluation.evaluate import calculate_metrics
from evaluation.dataset import GAIATask


class MockLLMClient:
    """Mock LLM client returning configurable LLMResponse and recording prompts/parts."""

    def __init__(self, response_text: str = "42", finish_reason: str = "STOP"):
        self.model = "gemini-3.5-flash-lite"
        self.temperature = None
        self.max_output_tokens = 2048
        self.thinking_level = "medium"
        self.response = LLMResponse(
            text=response_text,
            raw_text=response_text,
            finish_reason=finish_reason,
            input_tokens=25,
            output_tokens=5,
            thinking_tokens=50,
            total_tokens=80,
            response_id="mock-file-resp-id",
            model_version="gemini-3.5-flash-lite",
        )
        self.last_prompt = None
        self.last_attachment_parts = None

    def generate(self, prompt: str, attachment_parts=None) -> LLMResponse:
        self.last_prompt = prompt
        self.last_attachment_parts = attachment_parts
        return self.response


class MockSearchTool:
    """Mock Tavily search tool returning deterministic results."""

    def __init__(self, success: bool = True, num_results: int = 2):
        self.success = success
        self.num_results = num_results
        self.call_count = 0
        self.last_query = None

    def search(self, query: str) -> WebSearchResult:
        self.call_count += 1
        self.last_query = query
        if not self.success:
            return WebSearchResult(
                query=query,
                results=[],
                success=False,
                latency_seconds=0.05,
                error_type="SearchError",
                error_message="Mock search failure",
                call_count=1,
                provider="tavily",
            )
        items = [
            SearchResultItem(
                title=f"Result {i}",
                url=f"https://example.com/{i}",
                content=f"Search evidence snippet {i}",
                score=0.9 - (i * 0.1),
            )
            for i in range(self.num_results)
        ]
        return WebSearchResult(
            query=query,
            results=items,
            success=True,
            latency_seconds=0.1,
            call_count=1,
            provider="tavily",
        )


class TestFileToolProcessing(unittest.TestCase):
    """Unit tests for FileTool extraction across supported file types."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.tool = FileTool(max_text_chars=50000)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_text_file_processing(self):
        file_path = os.path.join(self.temp_dir.name, "sample.txt")
        content = "Hello, world! This is a test file for GAIA V2."
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)

        result = self.tool.process_file(file_path)
        self.assertTrue(result.success)
        self.assertEqual(result.content_mode, "text")
        self.assertEqual(result.processor, "plain_text")
        self.assertIn(content, result.text_content)
        self.assertFalse(result.content_truncated)
        self.assertEqual(result.original_content_length, len(content))
        self.assertEqual(result.provided_content_length, len(content))

    def test_json_and_csv_processing(self):
        csv_path = os.path.join(self.temp_dir.name, "data.csv")
        with open(csv_path, "w", encoding="utf-8") as f:
            f.write("id,name,score\n1,Alpha,95\n2,Beta,88\n")

        res_csv = self.tool.process_file(csv_path)
        self.assertTrue(res_csv.success)
        self.assertEqual(res_csv.content_mode, "text")
        self.assertEqual(res_csv.processor, "plain_text")
        self.assertIn("1,Alpha,95", res_csv.text_content)

        json_path = os.path.join(self.temp_dir.name, "data.json")
        with open(json_path, "w", encoding="utf-8") as f:
            f.write('{"key": "value", "count": 42}')

        res_json = self.tool.process_file(json_path)
        self.assertTrue(res_json.success)
        self.assertEqual(res_json.content_mode, "text")
        self.assertEqual(res_json.processor, "plain_text")
        self.assertIn('"key": "value"', res_json.text_content)

    def test_text_truncation_boundary(self):
        tool_small = FileTool(max_text_chars=100)
        file_path = os.path.join(self.temp_dir.name, "large.txt")
        large_content = "A" * 250
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(large_content)

        result = tool_small.process_file(file_path)
        self.assertTrue(result.success)
        self.assertTrue(result.content_truncated)
        self.assertEqual(result.original_content_length, 250)
        self.assertEqual(result.provided_content_length, 100)
        self.assertEqual(result.text_content, "A" * 100)
        self.assertEqual(len(result.text_content), 100)

    def test_python_file_no_execution_security(self):
        """CRITICAL SECURITY TEST: Ensure python files are treated as text and never executed."""
        marker_file = os.path.join(self.temp_dir.name, "MALICIOUS_EXECUTION_MARKER.txt")
        py_file = os.path.join(self.temp_dir.name, "script.py")

        # Malicious code that writes a file if executed via exec(), eval(), or subprocess
        malicious_code = f"""
import os
with open(r'{marker_file}', 'w') as f:
    f.write('EXPLOIT_EXECUTED')
answer = 42
"""
        with open(py_file, "w", encoding="utf-8") as f:
            f.write(malicious_code)

        result = self.tool.process_file(py_file)

        # 1. Result should succeed as text
        self.assertTrue(result.success)
        self.assertEqual(result.content_mode, "text")
        self.assertEqual(result.processor, "python_source")
        self.assertIn("answer = 42", result.text_content)

        # 2. Disk marker must NOT exist (strictly verifying no code execution occurred)
        self.assertFalse(
            os.path.exists(marker_file),
            "SECURITY VIOLATION: Python attachment was executed!",
        )

    def test_docx_processing(self):
        docx_path = os.path.join(self.temp_dir.name, "document.docx")
        doc = docx.Document()
        doc.add_heading("GAIA Research Report", level=1)
        doc.add_paragraph("This is an evaluation paragraph.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Header A"
        table.cell(0, 1).text = "Header B"
        table.cell(1, 0).text = "Value 1"
        table.cell(1, 1).text = "Value 2"
        doc.save(docx_path)

        result = self.tool.process_file(docx_path)
        self.assertTrue(result.success)
        self.assertEqual(result.content_mode, "text")
        self.assertEqual(result.processor, "python_docx")
        self.assertIn("GAIA Research Report", result.text_content)
        self.assertIn("This is an evaluation paragraph.", result.text_content)
        self.assertIn("Header A", result.text_content)
        self.assertIn("Value 2", result.text_content)

    def test_xlsx_processing(self):
        xlsx_path = os.path.join(self.temp_dir.name, "spreadsheet.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Financials"
        ws["A1"] = "Metric"
        ws["B1"] = "Q1"
        ws["A2"] = "Revenue"
        ws["B2"] = 1000000
        wb.save(xlsx_path)

        result = self.tool.process_file(xlsx_path)
        self.assertTrue(result.success)
        self.assertEqual(result.content_mode, "text")
        self.assertEqual(result.processor, "openpyxl")
        self.assertIn("Financials", result.text_content)
        self.assertIn("A1 | value=Metric", result.text_content)
        self.assertIn("B2 | value=1000000", result.text_content)

    @patch("openpyxl.load_workbook")
    def test_xls_rejected_as_unsupported_without_calling_openpyxl(self, mock_load_workbook):
        """Verify legacy binary .xls is rejected deterministically and openpyxl is never invoked."""
        xls_path = os.path.join(self.temp_dir.name, "legacy.xls")
        with open(xls_path, "wb") as f:
            f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

        result = self.tool.process_file(xls_path)

        # 1. Rejected deterministically
        self.assertFalse(result.success)
        self.assertEqual(result.error_type, "UnsupportedFileTypeError")
        self.assertEqual(result.file_extension, ".xls")
        self.assertIn("not supported in V2", result.error_message)
        self.assertIsNone(result.processor)

        # 2. openpyxl loader must never be called
        mock_load_workbook.assert_not_called()

    def test_pptx_processing(self):
        pptx_path = os.path.join(self.temp_dir.name, "presentation.pptx")
        prs = Presentation()
        blank_slide_layout = prs.slide_layouts[6]  # blank layout
        slide = prs.slides.add_slide(blank_slide_layout)
        tx_box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(2))
        tf = tx_box.text_frame
        tf.text = "Presentation Slide 1 Heading"
        p = tf.add_paragraph()
        p.text = "Bullet item detail"
        prs.save(pptx_path)

        result = self.tool.process_file(pptx_path)
        self.assertTrue(result.success)
        self.assertEqual(result.content_mode, "text")
        self.assertEqual(result.processor, "python_pptx")
        self.assertIn("Slide 1", result.text_content)
        self.assertIn("Presentation Slide 1 Heading", result.text_content)
        self.assertIn("Bullet item detail", result.text_content)

    def test_native_multimodal_dispatch(self):
        # Image PNG
        png_path = os.path.join(self.temp_dir.name, "image.png")
        dummy_png_bytes = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4"
        with open(png_path, "wb") as f:
            f.write(dummy_png_bytes)

        res_png = self.tool.process_file(png_path)
        self.assertTrue(res_png.success)
        self.assertEqual(res_png.content_mode, "native_multimodal")
        self.assertEqual(res_png.mime_type, "image/png")
        self.assertEqual(res_png.native_bytes, dummy_png_bytes)
        self.assertEqual(res_png.processor, "gemini_multimodal_image")

        # Audio MP3
        mp3_path = os.path.join(self.temp_dir.name, "audio.mp3")
        dummy_mp3_bytes = b"ID3\x03\x00\x00\x00\x00\x00\x00mockaudiobytes"
        with open(mp3_path, "wb") as f:
            f.write(dummy_mp3_bytes)

        res_mp3 = self.tool.process_file(mp3_path)
        self.assertTrue(res_mp3.success)
        self.assertEqual(res_mp3.content_mode, "native_multimodal")
        self.assertEqual(res_mp3.mime_type, "audio/mp3")
        self.assertEqual(res_mp3.native_bytes, dummy_mp3_bytes)
        self.assertEqual(res_mp3.processor, "gemini_multimodal_audio")

        # PDF native multimodal
        pdf_path = os.path.join(self.temp_dir.name, "doc.pdf")
        dummy_pdf_bytes = b"%PDF-1.4\n%mock pdf header\n%%EOF"
        with open(pdf_path, "wb") as f:
            f.write(dummy_pdf_bytes)

        res_pdf = self.tool.process_file(pdf_path)
        self.assertTrue(res_pdf.success)
        self.assertEqual(res_pdf.content_mode, "native_multimodal")
        self.assertEqual(res_pdf.mime_type, "application/pdf")
        self.assertEqual(res_pdf.native_bytes, dummy_pdf_bytes)
        self.assertEqual(res_pdf.processor, "gemini_multimodal_pdf")

    def test_missing_file_handling(self):
        missing_path = os.path.join(self.temp_dir.name, "nonexistent.txt")
        result = self.tool.process_file(missing_path)

        self.assertFalse(result.success)
        self.assertEqual(result.error_type, "FileNotFoundError")
        self.assertIn("not found", result.error_message)

    def test_unsupported_file_extension(self):
        unsupported_path = os.path.join(self.temp_dir.name, "binary.exe")
        with open(unsupported_path, "wb") as f:
            f.write(b"MZ\x90\x00")

        result = self.tool.process_file(unsupported_path)
        self.assertFalse(result.success)
        self.assertEqual(result.error_type, "UnsupportedFileTypeError")
        self.assertIn("not supported in V2", result.error_message)


class TestGAIAFileAgent(unittest.TestCase):
    """Unit tests for GAIAFileAgent behavior, search invariance, and fallback logic."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.llm = MockLLMClient(response_text="Paris")
        self.search_tool = MockSearchTool(success=True)
        self.file_tool = FileTool()
        self.agent = GAIAFileAgent(
            llm_client=self.llm,
            search_tool=self.search_tool,
            file_tool=self.file_tool,
        )

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_single_shot_search_invariant_with_attachment(self):
        """CRITICAL ABLATION INVARIANT: Even when an attachment is processed, Tavily search must be called once."""
        txt_path = os.path.join(self.temp_dir.name, "sample.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Attachment context data.")

        question = "What is the capital of France?"
        result = self.agent.run(question, file_path=txt_path)

        # 1. Search must have been called exactly once
        self.assertEqual(self.search_tool.call_count, 1)
        self.assertEqual(self.search_tool.last_query, question)

        # 2. File must have been processed
        self.assertIsNotNone(result.file_result)
        self.assertTrue(result.file_result.success)

        # 3. Prompt version must be file-search-v1
        self.assertEqual(result.prompt_version, FILE_SEARCH_PROMPT_VERSION)
        self.assertEqual(result.primary_prompt_version, FILE_SEARCH_PROMPT_VERSION)
        self.assertFalse(result.file_fallback)

        # 4. LLM prompt must contain both web evidence and file content
        self.assertIn("WEB SEARCH EVIDENCE", self.llm.last_prompt)
        self.assertIn("ATTACHMENT EVIDENCE", self.llm.last_prompt)
        self.assertIn("Attachment context data.", self.llm.last_prompt)
        self.assertIn(question, self.llm.last_prompt)

    def test_no_attachment_falls_back_to_v1_prompt(self):
        """When no attachment is provided, agent produces standard V1 web-search prompt."""
        question = "What is the capital of France?"
        result = self.agent.run(question, file_path=None)

        self.assertEqual(self.search_tool.call_count, 1)
        self.assertIsNone(result.file_result)
        self.assertEqual(result.prompt_version, WEB_SEARCH_PROMPT_VERSION)
        self.assertIn("Web Search Evidence:", self.llm.last_prompt)
        self.assertNotIn("ATTACHMENT EVIDENCE", self.llm.last_prompt)

    def test_missing_file_fallback_to_web_search_prompt(self):
        """When file cannot be found, agent falls back to web-search-v1 prompt and sets file_fallback=True."""
        missing_path = os.path.join(self.temp_dir.name, "missing.txt")
        question = "What is the capital of France?"
        result = self.agent.run(question, file_path=missing_path)

        self.assertEqual(self.search_tool.call_count, 1)
        self.assertIsNotNone(result.file_result)
        self.assertFalse(result.file_result.success)
        self.assertTrue(result.file_fallback)
        self.assertEqual(result.prompt_version, WEB_SEARCH_PROMPT_VERSION)
        self.assertEqual(result.primary_prompt_version, FILE_SEARCH_PROMPT_VERSION)
        self.assertEqual(result.fallback_prompt_version, WEB_SEARCH_PROMPT_VERSION)

    def test_native_multimodal_dispatch_passes_parts_to_llm(self):
        """When an image attachment is provided, Gemini client receives a multimodal Part object."""
        png_path = os.path.join(self.temp_dir.name, "chart.png")
        png_bytes = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01\x08\x06\x00\x00\x00\x1f\x15c4"
        with open(png_path, "wb") as f:
            f.write(png_bytes)

        question = "What is shown in this chart?"
        result = self.agent.run(question, file_path=png_path)

        self.assertIsNotNone(result.file_result)
        self.assertEqual(result.file_result.content_mode, "native_multimodal")
        self.assertIsNotNone(self.llm.last_attachment_parts)
        self.assertEqual(len(self.llm.last_attachment_parts), 1)

    @patch("openpyxl.load_workbook")
    def test_xls_attachment_triggers_file_fallback(self, mock_load_workbook):
        """When a .xls file is encountered, agent triggers file fallback without invoking openpyxl."""
        xls_path = os.path.join(self.temp_dir.name, "data.xls")
        with open(xls_path, "wb") as f:
            f.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

        question = "What is the sum in data.xls?"
        result = self.agent.run(question, file_path=xls_path)

        # 1. Search must be called once
        self.assertEqual(self.search_tool.call_count, 1)

        # 2. openpyxl loader must never be called
        mock_load_workbook.assert_not_called()

        # 3. File result rejected as UnsupportedFileTypeError
        self.assertIsNotNone(result.file_result)
        self.assertFalse(result.file_result.success)
        self.assertEqual(result.file_result.error_type, "UnsupportedFileTypeError")
        self.assertTrue(result.file_fallback)

        # 4. Falls back to web-search-v1 prompt
        self.assertEqual(result.prompt_version, WEB_SEARCH_PROMPT_VERSION)
        self.assertEqual(result.primary_prompt_version, FILE_SEARCH_PROMPT_VERSION)
        self.assertEqual(result.fallback_prompt_version, WEB_SEARCH_PROMPT_VERSION)
        self.assertIn("Web Search Evidence:", self.llm.last_prompt)
        self.assertNotIn("ATTACHMENT EVIDENCE", self.llm.last_prompt)


class TestV2RunnerAndEvaluationPipeline(unittest.TestCase):
    """Tests for runner.py and evaluate.py V2 metric calculations."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_runner_v2_metadata_logging(self):
        txt_path = os.path.join(self.temp_dir.name, "attachment.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Some attachment facts.")

        task = GAIATask(
            task_id="test-task-v2-001",
            question="What is the fact in the document?",
            level=1,
            final_answer="facts",
            file_name="attachment.txt",
            file_path=txt_path,
        )

        llm = MockLLMClient(response_text="facts")
        search = MockSearchTool(success=True)
        file_tool = FileTool()
        agent = GAIAFileAgent(llm_client=llm, search_tool=search, file_tool=file_tool)

        record = execute_task(
            task_id=task.task_id,
            question=task.question,
            level=task.level,
            file_name=task.file_name,
            file_path=txt_path,
            agent=agent,
            llm=llm,
            project_version="v2",
        )

        # Verify all V2 file fields exist in the execution record
        self.assertEqual(record["project_version"], "v2")
        self.assertTrue(record["file_enabled"])
        self.assertTrue(record["file_present"])
        self.assertEqual(record["file_extension"], ".txt")
        self.assertTrue(record["file_processing_attempted"])
        self.assertTrue(record["file_processing_success"])
        self.assertEqual(record["file_processor"], "plain_text")
        self.assertEqual(record["file_content_mode"], "text")
        self.assertFalse(record["file_content_truncated"])
        self.assertFalse(record["file_fallback"])
        self.assertTrue(record["search_enabled"])
        self.assertEqual(record["search_call_count"], 1)

    def test_runner_v2_records_xls_unsupported_and_fallback(self):
        xls_path = os.path.join(self.temp_dir.name, "old.xls")
        with open(xls_path, "wb") as f:
            f.write(b"\xd0\xcf\x11\xe0")

        task = GAIATask(
            task_id="test-task-v2-xls",
            question="Analyze this old spreadsheet",
            level=1,
            final_answer="answer",
            file_name="old.xls",
            file_path=xls_path,
        )

        llm = MockLLMClient(response_text="answer")
        search = MockSearchTool(success=True)
        file_tool = FileTool()
        agent = GAIAFileAgent(llm_client=llm, search_tool=search, file_tool=file_tool)

        record = execute_task(
            task_id=task.task_id,
            question=task.question,
            level=task.level,
            file_name=task.file_name,
            file_path=xls_path,
            agent=agent,
            llm=llm,
            project_version="v2",
        )

        self.assertTrue(record["file_enabled"])
        self.assertTrue(record["file_present"])
        self.assertEqual(record["file_extension"], ".xls")
        self.assertTrue(record["file_processing_attempted"])
        self.assertFalse(record["file_processing_success"])
        self.assertEqual(record["file_error_type"], "UnsupportedFileTypeError")
        self.assertTrue(record["file_fallback"])

    def test_calculate_metrics_aggregates_v2_file_statistics(self):
        mock_tasks = {
            "task-1": GAIATask("task-1", "Q1", 1, "Ans1", file_name="doc.txt"),
            "task-2": GAIATask("task-2", "Q2", 1, "Ans2", file_name="table.xlsx"),
            "task-3": GAIATask("task-3", "Q3", 1, "Ans3"),
        }

        predictions = [
            {
                "task_id": "task-1",
                "level": 1,
                "project_version": "v2",
                "request_success": True,
                "completion_success": True,
                "final_answer": "Ans1",
                "attachment_required": True,
                "search_enabled": True,
                "search_success": True,
                "search_call_count": 1,
                "file_enabled": True,
                "file_present": True,
                "file_extension": ".txt",
                "file_processing_attempted": True,
                "file_processing_success": True,
                "file_processing_latency_seconds": 0.05,
                "file_processor": "plain_text",
                "file_content_mode": "text",
                "file_content_truncated": False,
                "original_file_content_length": 50,
                "provided_file_content_length": 50,
                "file_fallback": False,
            },
            {
                "task_id": "task-2",
                "level": 1,
                "project_version": "v2",
                "request_success": True,
                "completion_success": True,
                "final_answer": "Wrong",
                "attachment_required": True,
                "search_enabled": True,
                "search_success": True,
                "search_call_count": 1,
                "file_enabled": True,
                "file_present": True,
                "file_extension": ".xlsx",
                "file_processing_attempted": True,
                "file_processing_success": True,
                "file_processing_latency_seconds": 0.12,
                "file_processor": "openpyxl",
                "file_content_mode": "text",
                "file_content_truncated": True,
                "original_file_content_length": 60000,
                "provided_file_content_length": 50000,
                "file_fallback": False,
            },
            {
                "task_id": "task-3",
                "level": 1,
                "project_version": "v2",
                "request_success": True,
                "completion_success": True,
                "final_answer": "Ans3",
                "attachment_required": False,
                "search_enabled": True,
                "search_success": True,
                "search_call_count": 1,
                "file_enabled": True,
                "file_present": False,
                "file_extension": None,
                "file_processing_attempted": False,
                "file_processing_success": False,
                "file_fallback": False,
            },
        ]

        eval_result = calculate_metrics(
            predictions=predictions,
            tasks_by_id=mock_tasks,
            level=1,
            project_version="v2",
        )
        summary = eval_result["summary"]

        self.assertEqual(summary["project_version"], "v2")
        self.assertEqual(summary["total_tasks"], 3)
        self.assertEqual(summary["correct_tasks"], 2)
        self.assertEqual(summary["accuracy"], 0.6667)
        self.assertTrue(summary["file_enabled"])
        self.assertEqual(summary["file_processing_attempts"], 2)
        self.assertEqual(summary["successful_file_processing_count"], 2)
        self.assertEqual(summary["file_processing_success_rate"], 1.0)
        self.assertEqual(summary["file_content_truncated_count"], 1)
        self.assertEqual(summary["file_fallback_count"], 0)

        # Check extension breakdown
        ext_breakdown = summary["extension_breakdown"]
        self.assertIn(".txt", ext_breakdown)
        self.assertIn(".xlsx", ext_breakdown)
        self.assertEqual(ext_breakdown[".txt"]["total"], 1)
        self.assertEqual(ext_breakdown[".txt"]["correct"], 1)
        self.assertEqual(ext_breakdown[".xlsx"]["total"], 1)
        self.assertEqual(ext_breakdown[".xlsx"]["correct"], 0)


if __name__ == "__main__":
    unittest.main()

