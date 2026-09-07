"""Provider-independent FileTool and normalized FileResult for GAIA V2.

Processes local GAIA task attachments (text, python, docx, xlsx, pptx, image, audio, pdf)
into normalized text or native multimodal parts while enforcing deterministic context limits
and strict execution boundaries (no code execution, no OCR).
"""

import os
import time
from dataclasses import dataclass, field
from typing import Optional, Dict, Any, List


@dataclass
class FileResult:
    """Encapsulates the normalized result of local file inspection."""
    file_name: str
    file_extension: str
    success: bool

    content_mode: Optional[str] = None      # 'text' or 'native_multimodal'
    text_content: Optional[str] = None      # Extracted text representation
    native_file_path: Optional[str] = None  # Absolute or relative path to local file
    native_bytes: Optional[bytes] = field(default=None, repr=False) # Raw bytes for multimodal API
    mime_type: Optional[str] = None         # MIME type for multimodal parts
    processor: Optional[str] = None         # Identifier of handler/library used

    latency_seconds: float = 0.0

    content_truncated: bool = False
    original_content_length: int = 0
    provided_content_length: int = 0

    error_type: Optional[str] = None
    error_message: Optional[str] = None

    metadata: Dict[str, Any] = field(default_factory=dict)

    def to_dict(self) -> Dict[str, Any]:
        """Returns safe, serializable metadata dictionary excluding binary content."""
        return {
            "file_name": self.file_name,
            "file_extension": self.file_extension,
            "success": self.success,
            "content_mode": self.content_mode,
            "processor": self.processor,
            "latency_seconds": self.latency_seconds,
            "content_truncated": self.content_truncated,
            "original_content_length": self.original_content_length,
            "provided_content_length": self.provided_content_length,
            "error_type": self.error_type,
            "error_message": self.error_message,
            "metadata": self.metadata,
        }


class FileTool:
    """Deterministic local file processor for GAIA task attachments."""

    DEFAULT_MAX_TEXT_CHARS = 50000

    TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json"}
    PYTHON_EXTENSIONS = {".py"}
    DOCX_EXTENSIONS = {".docx"}
    XLSX_EXTENSIONS = {".xlsx"}
    PPTX_EXTENSIONS = {".pptx"}
    IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
    AUDIO_EXTENSIONS = {".mp3"}
    PDF_EXTENSIONS = {".pdf"}

    MIME_TYPES = {
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".mp3": "audio/mp3",
        ".pdf": "application/pdf",
    }

    def __init__(self, max_text_chars: int = DEFAULT_MAX_TEXT_CHARS, native_pdf: bool = True):
        self.max_text_chars = max_text_chars
        self.native_pdf = native_pdf

    def process(self, file_path: str) -> FileResult:
        """Processes a local file into a normalized FileResult without code execution or OCR."""
        start_time = time.time()
        file_name = os.path.basename(file_path) if file_path else ""
        _, ext = os.path.splitext(file_name)
        ext = ext.lower()

        if not file_path or not os.path.exists(file_path):
            latency = round(time.time() - start_time, 3)
            return FileResult(
                file_name=file_name,
                file_extension=ext,
                success=False,
                latency_seconds=latency,
                error_type="FileNotFoundError",
                error_message=f"Attachment file not found at path: {file_path}",
            )

        try:
            if ext in self.TEXT_EXTENSIONS:
                result = self._read_text(file_path, file_name, ext)
            elif ext in self.PYTHON_EXTENSIONS:
                # Strictly read source text only. Never execute or import.
                result = self._read_python_source(file_path, file_name, ext)
            elif ext in self.DOCX_EXTENSIONS:
                result = self._read_docx(file_path, file_name, ext)
            elif ext in self.XLSX_EXTENSIONS:
                result = self._read_xlsx(file_path, file_name, ext)
            elif ext in self.PPTX_EXTENSIONS:
                result = self._read_pptx(file_path, file_name, ext)
            elif ext in self.IMAGE_EXTENSIONS:
                result = self._read_multimodal(file_path, file_name, ext, "gemini_multimodal_image")
            elif ext in self.AUDIO_EXTENSIONS:
                result = self._read_multimodal(file_path, file_name, ext, "gemini_multimodal_audio")
            elif ext in self.PDF_EXTENSIONS:
                if self.native_pdf:
                    result = self._read_multimodal(file_path, file_name, ext, "gemini_multimodal_pdf")
                else:
                    result = self._read_pdf_text(file_path, file_name, ext)
            else:
                latency = round(time.time() - start_time, 3)
                return FileResult(
                    file_name=file_name,
                    file_extension=ext,
                    success=False,
                    latency_seconds=latency,
                    error_type="UnsupportedFileTypeError",
                    error_message=f"File extension '{ext}' is not supported in V2.",
                )

            # Apply deterministic truncation for text content
            if result.text_content is not None:
                orig_len = len(result.text_content)
                result.original_content_length = orig_len
                if orig_len > self.max_text_chars:
                    result.text_content = result.text_content[:self.max_text_chars]
                    result.content_truncated = True
                    result.provided_content_length = self.max_text_chars
                else:
                    result.content_truncated = False
                    result.provided_content_length = orig_len

            result.latency_seconds = round(time.time() - start_time, 3)
            return result

        except Exception as e:
            latency = round(time.time() - start_time, 3)
            return FileResult(
                file_name=file_name,
                file_extension=ext,
                success=False,
                latency_seconds=latency,
                error_type=type(e).__name__,
                error_message=str(e),
            )

    def _read_text(self, path: str, name: str, ext: str) -> FileResult:
        with open(path, "rb") as f:
            raw_bytes = f.read()

        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            try:
                text = raw_bytes.decode("latin-1")
            except Exception:
                text = raw_bytes.decode("utf-8", errors="replace")

        return FileResult(
            file_name=name,
            file_extension=ext,
            success=True,
            content_mode="text",
            text_content=text,
            processor="plain_text",
            metadata={"byte_size": len(raw_bytes)},
        )

    def _read_python_source(self, path: str, name: str, ext: str) -> FileResult:
        """Reads Python source code strictly as text. NEVER executes or imports code."""
        with open(path, "rb") as f:
            raw_bytes = f.read()

        try:
            text = raw_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = raw_bytes.decode("utf-8", errors="replace")

        header = f"# Python Source File: {name}\n# (Static source representation - not executed)\n\n"
        return FileResult(
            file_name=name,
            file_extension=ext,
            success=True,
            content_mode="text",
            text_content=header + text,
            processor="python_source",
            metadata={"byte_size": len(raw_bytes), "executed": False},
        )

    def _read_docx(self, path: str, name: str, ext: str) -> FileResult:
        import docx

        doc = docx.Document(path)
        blocks: List[str] = []

        para_idx = 1
        for p in doc.paragraphs:
            txt = p.text.strip()
            if txt:
                blocks.append(f"[Paragraph {para_idx}] {txt}")
                para_idx += 1

        for t_idx, table in enumerate(doc.tables, start=1):
            table_lines: List[str] = [f"[Table {t_idx}]"]
            for r_idx, row in enumerate(table.rows, start=1):
                cells_txt = [c.text.strip().replace("\n", " ") for c in row.cells]
                table_lines.append(f"Row {r_idx}: " + " | ".join(cells_txt))
            blocks.append("\n".join(table_lines))

        full_text = "\n\n".join(blocks)
        return FileResult(
            file_name=name,
            file_extension=ext,
            success=True,
            content_mode="text",
            text_content=full_text,
            processor="python_docx",
            metadata={"paragraphs_count": para_idx - 1, "tables_count": len(doc.tables)},
        )

    def _read_xlsx(self, path: str, name: str, ext: str) -> FileResult:
        import openpyxl

        wb = openpyxl.load_workbook(path, data_only=True)
        sheet_blocks: List[str] = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            lines: List[str] = [f"Sheet: {sheet_name}"]

            # Collect merged cell ranges
            merged_info = [str(rng) for rng in ws.merged_cells.ranges]
            if merged_info:
                lines.append(f"Merged Cells: {', '.join(merged_info)}")

            # Iterate rows and cells
            row_count = 0
            for row in ws.iter_rows():
                cell_parts = []
                for cell in row:
                    val = cell.value
                    if val is not None and str(val).strip() != "":
                        coord = cell.coordinate
                        part = f"{coord} | value={val}"

                        # Cell style metadata
                        fill = cell.fill
                        if fill and hasattr(fill, "fill_type") and fill.fill_type:
                            fg = getattr(fill, "fgColor", None)
                            rgb = getattr(fg, "rgb", None) if fg else None
                            if rgb and str(rgb) not in ("00000000", "0", "None"):
                                part += f" | fill={rgb}"

                        font = cell.font
                        if font:
                            fc = getattr(font, "color", None)
                            font_rgb = getattr(fc, "rgb", None) if fc else None
                            if font_rgb and str(font_rgb) not in ("00000000", "0", "None"):
                                part += f" | font_color={font_rgb}"
                            if getattr(font, "bold", False):
                                part += " | bold=True"

                        if cell.number_format and cell.number_format != "General":
                            part += f" | format={cell.number_format}"

                        cell_parts.append(part)

                if cell_parts:
                    row_count += 1
                    lines.append(" ; ".join(cell_parts))

            sheet_blocks.append("\n".join(lines))

        return FileResult(
            file_name=name,
            file_extension=ext,
            success=True,
            content_mode="text",
            text_content="\n\n".join(sheet_blocks),
            processor="openpyxl",
            metadata={"sheets": wb.sheetnames},
        )

    def _read_pptx(self, path: str, name: str, ext: str) -> FileResult:
        from pptx import Presentation

        prs = Presentation(path)
        slides_text: List[str] = []

        for s_idx, slide in enumerate(prs.slides, start=1):
            lines: List[str] = [f"[Slide {s_idx}]"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            lines.append(f"Text: {t}")
                if shape.has_table:
                    lines.append("[Table]")
                    for r_idx, row in enumerate(shape.table.rows, start=1):
                        row_vals = [c.text.strip().replace("\n", " ") for c in row.cells]
                        lines.append(f"Row {r_idx}: " + " | ".join(row_vals))

            slides_text.append("\n".join(lines))

        return FileResult(
            file_name=name,
            file_extension=ext,
            success=True,
            content_mode="text",
            text_content="\n\n".join(slides_text),
            processor="python_pptx",
            metadata={"slides_count": len(prs.slides)},
        )

    def _read_multimodal(self, path: str, name: str, ext: str, processor: str) -> FileResult:
        """Reads binary file for native multimodal Gemini inference."""
        with open(path, "rb") as f:
            data = f.read()

        mime = self.MIME_TYPES.get(ext, "application/octet-stream")
        return FileResult(
            file_name=name,
            file_extension=ext,
            success=True,
            content_mode="native_multimodal",
            text_content=None,
            native_file_path=path,
            native_bytes=data,
            mime_type=mime,
            processor=processor,
            metadata={"byte_size": len(data), "mime_type": mime},
        )

    def _read_pdf_text(self, path: str, name: str, ext: str) -> FileResult:
        import pypdf

        reader = pypdf.PdfReader(path)
        pages_text: List[str] = []
        for p_idx, page in enumerate(reader.pages, start=1):
            txt = (page.extract_text() or "").strip()
            if txt:
                pages_text.append(f"[Page {p_idx}]\n{txt}")

        return FileResult(
            file_name=name,
            file_extension=ext,
            success=True,
            content_mode="text",
            text_content="\n\n".join(pages_text),
            processor="pypdf",
            metadata={"pages_count": len(reader.pages)},
        )

    # Convenience alias for tests and runners
    process_file = process

